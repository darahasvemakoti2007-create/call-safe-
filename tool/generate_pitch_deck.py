import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette (Modern Dark Cyber-Security Theme)
    BG_COLOR = RGBColor(11, 19, 43)        # #0B132B (Deep Navy)
    CARD_BG = RGBColor(28, 37, 65)         # #1C2541 (Dark Slate Blue)
    CARD_BORDER = RGBColor(58, 80, 107)    # #3A506B
    CYAN_ACCENT = RGBColor(56, 189, 248)   # #38BDF8 (Vibrant Cyan)
    GREEN_ACCENT = RGBColor(52, 211, 153)  # #34D399 (Emerald Green)
    ORANGE_ACCENT = RGBColor(251, 146, 60) # #FB923C (Warm Orange)
    RED_ACCENT = RGBColor(244, 63, 94)     # #F43F5E (Alert Rose)
    WHITE_TEXT = RGBColor(248, 250, 252)   # #F8FAFC
    MUTED_TEXT = RGBColor(148, 163, 184)   # #94A3B8

    def add_background(slide):
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5)
        )
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category="GUARDCALL — HACKATHON PITCH", subtitle_text=None):
        # Category / Kicker
        kicker_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_k = kicker_box.text_frame
        tf_k.word_wrap = True
        tf_k.margin_left = tf_k.margin_top = tf_k.margin_right = tf_k.margin_bottom = 0
        p_k = tf_k.paragraphs[0]
        p_k.text = category.upper()
        p_k.font.size = Pt(11)
        p_k.font.bold = True
        p_k.font.color.rgb = CYAN_ACCENT

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.7))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = WHITE_TEXT

        if subtitle_text:
            sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(11.7), Inches(0.4))
            tf_s = sub_box.text_frame
            tf_s.word_wrap = True
            tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = 0
            p_s = tf_s.paragraphs[0]
            p_s.text = subtitle_text
            p_s.font.size = Pt(14)
            p_s.font.color.rgb = MUTED_TEXT

    def add_card(slide, left, top, width, height, title="", items=None, accent_color=CYAN_ACCENT):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1.5)

        tb = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        if title:
            p0 = tf.paragraphs[0]
            p0.text = title
            p0.font.size = Pt(16)
            p0.font.bold = True
            p0.font.color.rgb = accent_color
            p0.space_after = Pt(10)

        if items:
            start_idx = 0 if not title else 1
            for i, item in enumerate(items):
                p = tf.add_paragraph() if (title or i > 0) else tf.paragraphs[0]
                p.text = item
                p.font.size = Pt(12)
                p.font.color.rgb = WHITE_TEXT
                p.space_after = Pt(8)
        return card

    def add_notes(slide, notes_text):
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = notes_text

    # ==========================================
    # SLIDE 1: Title & Hook
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    add_background(s1)

    # Title Card
    t_box = s1.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(7.5), Inches(4.5))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "GUARDCALL"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = CYAN_ACCENT

    p2 = tf1.add_paragraph()
    p2.text = "AI-Powered Phone Scam Shield"
    p2.font.size = Pt(28)
    p2.font.bold = True
    p2.font.color.rgb = WHITE_TEXT
    p2.space_before = Pt(8)
    p2.space_after = Pt(18)

    p3 = tf1.add_paragraph()
    p3.text = "Real-time conversational manipulation defense for vulnerable seniors — engineered privacy-first with zero raw audio retention and proactive family escalation."
    p3.font.size = Pt(15)
    p3.font.color.rgb = MUTED_TEXT
    p3.space_after = Pt(24)

    # Key Highlights Badges
    badges = [
        "🛡️ Zero-Audio & Zero-Transcript Privacy",
        "🧠 Real-Time Psychological Pattern Detection",
        "👨‍👩‍👧 Proactive Family Circle Escalation",
        "⚡ 100% Deterministic Offline Demo Center"
    ]
    for b in badges:
        pb = tf1.add_paragraph()
        pb.text = b
        pb.font.size = Pt(13)
        pb.font.bold = True
        pb.font.color.rgb = GREEN_ACCENT
        pb.space_after = Pt(6)

    # Screenshot on right
    if os.path.exists("step19_home_ready.png"):
        s1.shapes.add_picture("step19_home_ready.png", Inches(8.8), Inches(1.0), width=Inches(3.6))

    add_notes(s1, "Every 15 seconds, a senior is targeted by a phone scammer. Traditional spam blockers check caller IDs, but scammers spoof numbers in seconds. We built GuardCall: the first real-time conversational scam shield that protects seniors while preserving 100% privacy.")

    # ==========================================
    # SLIDE 2: The Problem
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_background(s2)
    add_header(s2, "The Problem: Sophisticated Psychological Coercion", "PROBLEM ANALYSIS",
               "Scammers don't hack devices—they hack human psychology through fear, authority, and urgency.")

    add_card(s2, Inches(0.8), Inches(2.2), Inches(3.6), Inches(4.5),
             "🚨 $10B+ Financial Drain",
             [
                 "• Impersonation & Digital Arrest scams cause catastrophic financial loss globally.",
                 "• Senior citizens and elderly individuals living alone are prime targets.",
                 "• Life savings are wiped out within a single 15-minute phone call."
             ], RED_ACCENT)

    add_card(s2, Inches(4.8), Inches(2.2), Inches(3.6), Inches(4.5),
             "🧠 Psychological Exploitation",
             [
                 "• Scammers engineer acute panic: 'Your account will be suspended in 10 minutes!'",
                 "• Fake Authority: Impersonating police, central banks, customs, and telecom officials.",
                 "• High cognitive overload renders traditional advice ineffective in the moment."
             ], ORANGE_ACCENT)

    add_card(s2, Inches(8.8), Inches(2.2), Inches(3.6), Inches(4.5),
             "🔒 Enforced Social Isolation",
             [
                 "• Scammers enforce secrecy: 'Do not tell your children or you will face immediate arrest.'",
                 "• The victim is kept on the line continuously to block second opinions.",
                 "• Seniors face high-stakes intimidation completely alone."
             ], CYAN_ACCENT)

    add_notes(s2, "Scammers no longer rely on simple phishing links. They call victims directly, impersonate law enforcement, and induce severe panic. The senior is trapped alone on the call, isolated from their loved ones, until their money is transferred.")

    # ==========================================
    # SLIDE 3: Why Existing Defenses Fail
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_background(s3)
    add_header(s3, "Why Existing Defenses Fall Short", "DEFENSE GAP",
               "Static phone number lookups and cloud-recording architectures cannot solve live phone fraud.")

    add_card(s3, Inches(0.8), Inches(2.2), Inches(3.6), Inches(4.5),
             "❌ Static Caller ID Blacklists",
             [
                 "• Apps like Truecaller rely on historical crowdsourced reports.",
                 "• Useless against VoIP spoofing, SIM swaps, and freshly purchased burner numbers.",
                 "• The first hundreds of victims are already defrauded before a number is flagged."
             ], RED_ACCENT)

    add_card(s3, Inches(4.8), Inches(2.2), Inches(3.6), Inches(4.5),
             "❌ Post-Mortem Bank Alerts",
             [
                 "• SMS security alerts arrive after the transaction is executed.",
                 "• Victims are pressured to bypass 2FA by reading out OTPs under duress.",
                 "• Reactive transaction freezes are too late to stop instant wire transfers."
             ], ORANGE_ACCENT)

    add_card(s3, Inches(8.8), Inches(2.2), Inches(3.6), Inches(4.5),
             "❌ Cloud Recording Violations",
             [
                 "• Cloud call recorders violate wiretap and telecommunications privacy laws.",
                 "• Uploading audio or transcripts creates extreme liability and user mistrust.",
                 "• Seniors refuse intrusive surveillance software that monitors personal life."
             ], CYAN_ACCENT)

    add_notes(s3, "Number-lookup apps only flag scammers after victims have already lost money. And recording phone calls to cloud servers creates massive privacy violations. We need a solution that analyzes conversational intent in real time without recording private audio.")

    # ==========================================
    # SLIDE 4: Proposed Solution
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_background(s4)
    add_header(s4, "The Solution: Behavioral Intent Recognition", "PROPOSED SOLUTION",
               "Shifting scam defense from 'Who is calling?' to 'What are they psychologically attempting to do?'")

    add_card(s4, Inches(0.8), Inches(2.2), Inches(7.5), Inches(4.5),
             "🛡️ GuardCall: Real-Time Protective Shield",
             [
                 "1. Behavioral Signal Detection: Identifies manipulation markers (urgency, legal threats, authority impersonation, credential harvesting) in real time.",
                 "2. Explainable Threat Scoring: Categorizes risk into 4 clear tiers (Safe, Suspicious, High Risk, Critical Threat) with plain-language rationales.",
                 "3. Proactive Senior HUD: Interrupts high-pressure calls with high-contrast, accessible warnings giving simple actions: 'END CALL' or 'ALERT FAMILY'.",
                 "4. Family Circle Escalation: Relays structured threat telemetry directly to designated caregivers without exposing the elder's private conversation.",
                 "5. False-Positive Immunity: Differentiates legitimate banking advisories from coercive phishing."
             ], GREEN_ACCENT)

    if os.path.exists("sample_alert_detail.png"):
        s4.shapes.add_picture("sample_alert_detail.png", Inches(8.8), Inches(1.8), width=Inches(3.7))

    add_notes(s4, "GuardCall acts like an intelligent guardian sitting right beside the senior. It listens to conversational tactics—urgency, threats, OTP demands—and displays clear, easy-to-read interventions, allowing the elder to end the call or immediately escalate to their family.")

    # ==========================================
    # SLIDE 5: Architecture & Detection Pipeline
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_background(s5)
    add_header(s5, "End-to-End System Architecture", "SYSTEM ARCHITECTURE",
               "Transient in-memory speech evaluation orchestrated with zero disk caching.")

    add_card(s5, Inches(0.8), Inches(2.2), Inches(3.6), Inches(2.1),
             "1. Audio Ingestion",
             [
                 "• Transient in-memory audio buffers.",
                 "• Zero audio files stored on device disk.",
                 "• Compliant with OS security boundaries."
             ], CYAN_ACCENT)

    add_card(s5, Inches(4.8), Inches(2.2), Inches(3.6), Inches(2.1),
             "2. Transient Speech-to-Text",
             [
                 "• Abstracted SpeechToTextService layer.",
                 "• Converts dialogue to short conversational chunks.",
                 "• Flushed immediately from RAM."
             ], CYAN_ACCENT)

    add_card(s5, Inches(8.8), Inches(2.2), Inches(3.6), Inches(2.1),
             "3. Detection Engine",
             [
                 "• Multi-turn compounding heuristic analysis.",
                 "• Context-aware banking filters.",
                 "• Sequence escalation tracking."
             ], CYAN_ACCENT)

    add_card(s5, Inches(0.8), Inches(4.6), Inches(5.6), Inches(2.1),
             "4. Harmonized Risk Tiering (0-100)",
             [
                 "• 0–39: Safe (No action)",
                 "• 40–70: Suspicious (Elevated awareness)",
                 "• 71–84: High Risk (Senior Warning HUD popup)",
                 "• 85–100: Critical Threat (Immediate escalation & siren)"
             ], ORANGE_ACCENT)

    add_card(s5, Inches(6.8), Inches(4.6), Inches(5.6), Inches(2.1),
             "5. Dual-Action Delivery",
             [
                 "• Senior Intervention: High-contrast HUD dialog with 'End Call'.",
                 "• Caregiver Circle: Relays structured threat metadata to family.",
                 "• Zero personal transcript shared with anyone."
             ], GREEN_ACCENT)

    add_notes(s5, "Our pipeline feeds transient audio chunks through speech recognition directly into a multi-turn heuristic engine. When compounding threats like urgency and OTP requests occur in sequence, the risk score surges past threshold, instantly triggering intervention dialogues.")

    # ==========================================
    # SLIDE 6: Live Demo & Scenarios
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_background(s6)
    add_header(s6, "Deterministic Demo Center & Scenarios", "LIVE DEMO READINESS",
               "100% offline, zero-network dependency built for reliable hackathon evaluation.")

    add_card(s6, Inches(0.8), Inches(2.2), Inches(7.5), Inches(4.5),
             "🧪 5 Scripted Multi-Turn Scenarios",
             [
                 "1. Fake KYC Update (Banking Impersonation):",
                 "   Escalates urgency ('Account blocked in 10 mins') ➔ OTP demand ➔ Surges to 75 / 100 High Risk.",
                 "2. Digital Arrest Threat (Law Enforcement):",
                 "   Fake police/customs intimidation claiming illegal parcel ➔ Surges to 90 / 100 Critical Threat.",
                 "3. Grandchild Emergency (Distress Extortion):",
                 "   Emotional panic claiming hospital emergency requiring urgent UPI transfer ➔ High Risk.",
                 "4. Legitimate Bank Alert (BENCHMARK):",
                 "   Official bank advisory: 'We never ask for OTP/PIN' ➔ Stays firmly at 0 / 100 Safe (Zero false positive).",
                 "5. Hospital Appointment Reschedule: Routine clinical call ➔ 0 / 100 Safe."
             ], CYAN_ACCENT)

    if os.path.exists("sample_2_scenarios.png"):
        s6.shapes.add_picture("sample_2_scenarios.png", Inches(8.8), Inches(1.8), width=Inches(3.7))

    add_notes(s6, "To prove reliability, we built a deterministic Demo Center with 5 diverse scenarios. In the KYC demo, notice how caller urgency elevates the risk to 75, triggering the intervention dialog. And crucially, legitimate bank calls remain at 0/100, proving zero false positives.")

    # ==========================================
    # SLIDE 7: Privacy by Design & Caregiver Boundary
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_background(s7)
    add_header(s7, "Privacy-by-Design: Zero Audio, Zero Transcripts", "PRIVACY & ETHICS",
               "Protecting seniors from fraud without sacrificing their personal dignity or privacy.")

    add_card(s7, Inches(0.8), Inches(2.2), Inches(5.6), Inches(4.5),
             "🔒 What Stays 100% Private to Senior",
             [
                 "❌ ZERO Audio Recorded or Shared:",
                 "   Microphone audio is never recorded, saved to disk, or streamed to caregivers.",
                 "",
                 "❌ ZERO Conversation Transcripts:",
                 "   Spoken words and sensitive dialogues are never saved or delivered in notifications.",
                 "",
                 "❌ ZERO PII / Credential Storage:",
                 "   OTPs, bank names, and passwords spoken during calls are never persisted."
             ], RED_ACCENT)

    add_card(s7, Inches(6.8), Inches(2.2), Inches(5.6), Inches(4.5),
             "🛡️ What Caregivers Receive (Telemetry Only)",
             [
                 "✅ Mathematical Threat Score:",
                 "   e.g. Risk Score: 75/100 (High Risk Threat).",
                 "",
                 "✅ Behavioral Manipulation Tags:",
                 "   e.g. ['Authority Impersonation', 'Urgency', 'OTP Demand'].",
                 "",
                 "✅ Senior Action Taken:",
                 "   e.g. 'Senior selected: Alert Family / End Call'.",
                 "",
                 "✅ Mutual Dual-Consent Authorization:",
                 "   Strictly requires mutual invite acceptance."
             ], GREEN_ACCENT)

    add_notes(s7, "Privacy is our core architectural pillar. Caregivers never hear the conversation or read personal transcripts. They receive only structured threat metadata: the risk score, the detected signals, and the action taken.")

    # ==========================================
    # SLIDE 8: Tech Stack & Quality
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_background(s8)
    add_header(s8, "Technical Stack & Production Hardening", "ENGINEERING RIGOR",
               "Modular mobile client architecture with 223 passing automated tests.")

    add_card(s8, Inches(0.8), Inches(2.2), Inches(5.6), Inches(2.1),
             "📱 Mobile Client (Frontend)",
             [
                 "• Flutter 3.x / Dart 3.x (Material 3)",
                 "• Custom Accessible Primitives (GlassCard, RiskBadge)",
                 "• Native Reduce Motion support & large touch targets"
             ], CYAN_ACCENT)

    add_card(s8, Inches(6.8), Inches(2.2), Inches(5.6), Inches(2.1),
             "🧠 Detection Engine (3 Modes)",
             [
                 "• Mode 1: In-memory Heuristic Mock Engine (Offline demos)",
                 "• Mode 2: Edge Function Proxy (Supabase Serverless)",
                 "• Mode 3: Direct Dev Mode (Gemini 1.5 Flash / GPT-4o-mini)"
             ], CYAN_ACCENT)

    add_card(s8, Inches(0.8), Inches(4.6), Inches(5.6), Inches(2.1),
             "☁️ Backend & Security",
             [
                 "• Supabase Auth & PostgreSQL database schema",
                 "• Row Level Security (RLS) enforcing mutual family ties",
                 "• Provider-neutral NotificationService architecture"
             ], CYAN_ACCENT)

    add_card(s8, Inches(6.8), Inches(4.6), Inches(5.6), Inches(2.1),
             "✅ Testing & Verification",
             [
                 "• 223 / 223 automated tests passing in Flutter test suite",
                 "• Zero dangerous runtime permissions in AndroidManifest",
                 "• AppLogger scrubbing credentials and sensitive tokens"
             ], GREEN_ACCENT)

    add_notes(s8, "Built with Flutter for smooth cross-platform delivery and Supabase for secure cloud synchronization. Our engine is completely modular, and our entire pipeline is backed by 223 passing automated tests.")

    # ==========================================
    # SLIDE 9: Roadmap & Future Vision
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_background(s9)
    add_header(s9, "Real-World Viability & Future Roadmap", "FUTURE ROADMAP",
               "Bridging hackathon prototype to nationwide telecom and consumer protection.")

    add_card(s9, Inches(0.8), Inches(2.2), Inches(3.6), Inches(4.5),
             "📍 Phase 1: Prototype (Today)",
             [
                 "• Standalone Flutter Android client.",
                 "• Deterministic 5-scenario Demo Center.",
                 "• Multi-turn heuristic scam engine.",
                 "• Privacy-first Family Circle escalation.",
                 "• 223/223 automated tests passing."
             ], CYAN_ACCENT)

    add_card(s9, Inches(4.8), Inches(2.2), Inches(3.6), Inches(4.5),
             "⚡ Phase 2: On-Device SLMs",
             [
                 "• Local Small Language Models (SLMs) via MediaPipe / ONNX.",
                 "• 100% offline, zero-latency on-device reasoning.",
                 "• Multi-lingual Indian dialect support (Hindi, Telugu, Tamil, Marathi).",
                 "• Voice biometrics for impersonation detection."
             ], ORANGE_ACCENT)

    add_card(s9, Inches(8.8), Inches(2.2), Inches(3.6), Inches(4.5),
             "🌐 Phase 3: Telco Partnerships",
             [
                 "• Integration with telecom carriers (Airtel, Jio, etc.) for native network-level shielding.",
                 "• Hardware-agnostic VoIP dialer integration.",
                 "• National Cybercrime Reporting Portal (NCRP) one-tap incident forwarding."
             ], GREEN_ACCENT)

    add_notes(s9, "We engineered GuardCall with accessibility at the center, including native reduce-motion toggles. While we demonstrate this today with controlled audio inputs, our roadmap includes deploying lightweight on-device SLMs and partnering directly with telecom carriers.")

    # ==========================================
    # SLIDE 10: Summary & Call to Action
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    add_background(s10)

    # Big Centered Summary Box
    s10_box = s10.shapes.add_textbox(Inches(1.5), Inches(0.8), Inches(10.3), Inches(2.2))
    tf10 = s10_box.text_frame
    tf10.word_wrap = True
    p10_1 = tf10.paragraphs[0]
    p10_1.text = "GuardCall: Restoring Confidence in Every Call"
    p10_1.font.size = Pt(32)
    p10_1.font.bold = True
    p10_1.font.color.rgb = CYAN_ACCENT
    p10_1.alignment = PP_ALIGN.CENTER

    p10_2 = tf10.add_paragraph()
    p10_2.text = "Seniors should never have to fear answering their phones, and families shouldn't have to worry."
    p10_2.font.size = Pt(16)
    p10_2.font.color.rgb = MUTED_TEXT
    p10_2.alignment = PP_ALIGN.CENTER
    p10_2.space_before = Pt(8)

    add_card(s10, Inches(1.2), Inches(3.0), Inches(3.3), Inches(2.6),
             "🛡️ Proactive Defense",
             [
                 "Detects psychological coercion dynamically before money is transferred."
             ], CYAN_ACCENT)

    add_card(s10, Inches(5.0), Inches(3.0), Inches(3.3), Inches(2.6),
             "🔒 Absolute Privacy",
             [
                 "Zero audio recorded, zero transcripts saved. 100% privacy by design."
             ], GREEN_ACCENT)

    add_card(s10, Inches(8.8), Inches(3.0), Inches(3.3), Inches(2.6),
             "👨‍👩‍👧 Family Connected",
             [
                 "Empowers elders with an instant safety bridge to their trusted caregivers."
             ], ORANGE_ACCENT)

    # Bottom Banner
    banner = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(6.0), Inches(10.3), Inches(0.9))
    banner.fill.solid()
    banner.fill.fore_color.rgb = CARD_BG
    banner.line.color.rgb = CYAN_ACCENT
    banner.line.width = Pt(1.5)

    b_tf = banner.text_frame
    b_tf.margin_left = b_tf.margin_top = b_tf.margin_right = b_tf.margin_bottom = 0
    p_b = b_tf.paragraphs[0]
    p_b.text = "✨ READY FOR LIVE DEMONSTRATION IN THE DEMO CENTER — THANK YOU! ✨"
    p_b.font.size = Pt(15)
    p_b.font.bold = True
    p_b.font.color.rgb = WHITE_TEXT
    p_b.alignment = PP_ALIGN.CENTER

    add_notes(s10, "GuardCall gives seniors the independence they deserve, and gives families peace of mind. Thank you, and we invite you to experience the live multi-turn simulation right now!")

    output_path = "GuardCall_Hackathon_Pitch.pptx"
    prs.save(output_path)
    print(f"Successfully generated PowerPoint presentation at: {os.path.abspath(output_path)}")

if __name__ == "__main__":
    create_deck()
